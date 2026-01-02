# 1/2/2026 
# TO DO: Add regex pattern for requirement IDs
# Added support for tracking headers to heading level 3..although I've seen them go all the way to 6.
# Added slight structure improvement in excel.


import re
from docx import Document
from openpyxl import Workbook

def extract_sentences_with_shall(docx_path):
    """Extracts sentences containing the word 'shall' from a Word document."""
    try:
        doc = Document(docx_path)
    except Exception as e:
        print("Failed to open input word document!")
        return []
    
    sentences_with_shall = []
    h1 = h2 = h3 = 0


    # Iterate through paragraphs in the document
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        style = paragraph.style.name

        # Section Numbers
        if style == "Heading 1":
            h1 += 1
            h2 = h3 = 0
            current_section = f"{h1}"

        elif style == "Heading 2":
            if h1 == 0:
                h1 = 1
            h2 += 1
            h3 = 0
            current_section = f"{h1}.{h2}"

        elif style == "Heading 3":
            if h1 == 0:
                h1 = 1
            if h2 == 0:
                h2 = 1
            h3 += 1
            current_section = f"{h1}.{h2}.{h3}"


        # Split the paragraph into sentences using regex
        sentences = re.split(r'(?<=[.!?]) +', paragraph.text)
        for sentence in sentences:
            if 'shall' in sentence.lower():  # Case-insensitive search for 'shall'
                sentences_with_shall.append({
                    "section": current_section, 
                    "description": sentence.strip()
                })
    
    return sentences_with_shall

def export_to_excel(req, output_excel_path):
    """Exports a list of sentences to an Excel file."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sentences with Shall"
    sheet.append(["REQ ID", "Section #", "Description")

    # Write sentences to individual rows
    for item in req:
        sheet.append(["TBD", req["section"], req["description"]])
    
    workbook.save(output_excel_path)

if __name__ == "__main__":
    input_docx = "input.docx"  # Replace with your Word document file path
    output_excel = "output.xlsx"  # Replace with your desired Excel file path

    # Extract sentences and export to Excel
    req = extract_sentences_with_shall(input_docx)
    if req:
        export_to_excel(req, output_excel)
        print(f"Exported {len(req)} sentences to '{output_excel}'.")
    else:
        print("No sentences containing 'shall' were found.")

